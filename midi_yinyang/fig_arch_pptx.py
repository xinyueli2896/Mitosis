"""The model figure as an editable PowerPoint slide (and a matplotlib
preview of the same drawing), following the layout of the hand-edited
deck Duet_larger_C_v23 and the notation of paper/method.tex:

  W^Q_s, W^K_s, W^V_s, W^O_s      role superscript, stream subscript
  K_intra / K_cross / K_frame     intra-stream, cross-stream, within-frame
  g^c, g^f                        the two scalar gates (Eq. gates)
  G^(x), G^(y), pi                per-stream routers and their softmax
  LN(h + o)                       post-LN: the add comes before the LN
  p(x_t | x_<t, y_<t)             decode conditionals in the setup's form

Panels: C (top) alternating-commit decoding; A (left) the transformer
block, input row below and predicted row above; B (right) who attends
to whom, from the model's own masks.

Usage:
  python fig_arch_pptx.py --out results/fig_model            # .pptx + .png preview
"""

import argparse
import os
import sys

import numpy as np
import torch

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

# ---------------------------------------------------------------- palette
BLUE, BLUE_L = '4589E3', 'CDDBF1'
RED, RED_L = 'BF2847', 'EDCCD3'
GOLD, GOLD_L = 'F0A500', 'FCE8B1'
INK, INK_2, GREY_L, PANEL = '303030', '737373', 'D9D9D9', 'EDECE9'
LINE_INK = '3A3A3A'                 # default colour of lines and arrows (fig_style.INK)
FONT = 'Times New Roman'
SLIDE_W, SLIDE_H = 14.79, 8.91


# ---------------------------------------------------------------- spec
class Spec:
    """A list of drawing primitives in inches, y down from the top."""

    def __init__(self):
        self.items = []

    def rect(self, x, y, w, h, fill='FFFFFF', line=INK, lw=0.75, dash=False, runs=None,
             size=10, align='c', color=INK, shape='rect', z=1):
        self.items.append(dict(k='rect', x=x, y=y, w=w, h=h, fill=fill, line=line, lw=lw,
                               dash=dash, runs=runs, size=size, align=align, color=color,
                               shape=shape, z=z))

    def text(self, x, y, w, h, runs, size=10, align='c', color=INK, bold=False, z=3,
             valign='m'):
        self.items.append(dict(k='text', x=x, y=y, w=w, h=h, runs=runs, size=size, align=align,
                               color=color, bold=bold, z=z, valign=valign))

    def line(self, x1, y1, x2, y2, color=LINE_INK, lw=0.75, dash=False, arrow=False, z=2):
        self.items.append(dict(k='line', x1=x1, y1=y1, x2=x2, y2=y2, color=color, lw=lw,
                               dash=dash, arrow=arrow, z=z))

    def curve(self, x1, y1, x2, y2, color=LINE_INK, lw=0.75, arrow=False, z=2, bend=0.5):
        """A vertical S-curve from (x1, y1) to (x2, y2): a cubic Bezier whose
        tangents are vertical at both ends; bend in (0, 1] sets how far
        the control points sit along the vertical span."""
        self.items.append(dict(k='curve', x1=x1, y1=y1, x2=x2, y2=y2, color=color, lw=lw,
                               arrow=arrow, z=z, bend=bend))

    def poly(self, points, fill, alpha=0.35, z=0):
        """A filled translucent polygon without outline; points in inches."""
        self.items.append(dict(k='poly', points=list(points), fill=fill, alpha=alpha, z=z))

    def image(self, x, y, w, h, path, z=5):
        """A picture (PNG with alpha) placed in the box (x, y, w, h)."""
        self.items.append(dict(k='image', x=x, y=y, w=w, h=h, path=os.path.abspath(path), z=z))

    def node(self, cx, cy, r=0.11, label='+', size=9):
        self.rect(cx - r, cy - r, 2 * r, 2 * r, fill='FFFFFF', line=INK, lw=0.75,
                  runs=[(label, '')], size=size, shape='ellipse', z=4)


# ---------------------------------------------------------------- provenance badges
ICON_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'figures', 'icons')
ICON = {'fire': (os.path.join(ICON_DIR, 'fire.png'), 503 / 585),
        'snow': (os.path.join(ICON_DIR, 'snowflake.png'), 548 / 626)}   # (path, w/h)


def badge(S, x_right, y_top, kind, h=0.15, z=6):
    """The provenance badge as pictures, right-aligned at x_right with its
    top at y_top: 'pre' = snowflake -> fire (initialised from the pretrained
    model, fine-tuned), 'new' = fire (trained from scratch). Returns its
    left edge."""
    gap, arrow_len = 0.035 * h / 0.15, 0.55 * h
    w_f = h * ICON['fire'][1]
    x = x_right - w_f
    S.image(x, y_top, w_f, h, ICON['fire'][0], z=z)
    if kind == 'pre':
        x -= gap + arrow_len
        # a thin shaft with a SMALL filled head (a polygon, so its size is
        # ours to set: the shape arrowheads of pptxgenjs cannot be scaled)
        ym = y_top + h / 2
        head = 0.22 * h
        S.line(x, ym, x + arrow_len - head * 0.8, ym, lw=0.8 * h / 0.15, z=z)
        S.poly([(x + arrow_len, ym), (x + arrow_len - head, ym - head * 0.5),
                (x + arrow_len - head, ym + head * 0.5)], fill=INK, alpha=1.0, z=z)
        w_s = h * ICON['snow'][1]
        x -= gap + w_s
        S.image(x, y_top, w_s, h, ICON['snow'][0], z=z)
    return x


# runs: list of (text, flags) with flags among 'i' (italic), 'sup', 'sub', 'b' (bold)
def m(*parts):
    """Shorthand: m('W', ('Q','sup'), ('x','sub')) -> runs; bare strings are italic."""
    runs = []
    for p in parts:
        if isinstance(p, str):
            runs.append((p, 'i'))
        else:
            t, f = p
            runs.append((t, f + ' i' if 'r' not in f else f.replace('r', '')))
    return runs


def plain(s):
    return [(s, '')]


def tok(sub):
    return sub


# ---------------------------------------------------------------- content
def build(T_show=4):
    S = Spec()
    # ------------------------------------------------ panel C: decode (top)
    S.text(4.0, 0.08, 4.0, 0.3, [('C', 'b'), ('  Alternating-commit decoding', '')], size=12,
           align='l')
    tw, tg = 0.3, 0.07

    def token_row(x, y, labels, kinds):
        """labels: runs per token; kinds: 'x' 'y' 'q' 'qm' (masked) 'qx' (holds x_t)."""
        xs = []
        for lab, kind in zip(labels, kinds):
            if lab is None:                       # ellipsis
                S.text(x, y, tw, tw, plain('…'), size=11)
                x += tw + tg; continue
            fill, line = {'x': (BLUE_L, BLUE), 'y': (RED_L, RED), 'q': (GOLD_L, GOLD),
                          'qm': ('FFFFFF', GOLD), 'qx': (BLUE_L, BLUE)}[kind]
            S.rect(x, y, tw, tw, fill=fill, line=line, lw=1.0, dash=(kind == 'qm'),
                   runs=m('x', ('t', 'sub')) if kind == 'qx' else None, size=8)
            S.text(x - 0.1, y + tw + 0.02, tw + 0.2, 0.2, lab, size=8)
            xs.append(x + tw / 2); x += tw + tg
        return xs, x

    prefix_labels = [m('x', ('1', 'sub')), m('y', ('1', 'sub')), None,
                     m('x', ('t−1', 'sub')), m('y', ('t−1', 'sub'))]
    prefix_kinds = ['x', 'y', None, 'x', 'y']
    steps = [
        (4.1, '1  draft the leader', ['qm', 'qm'],
         m('x', ('t', 'sub'), (' ∼ ', 'r'), 'p', ('(', 'r'), 'x', ('t', 'sub'), (' | ', 'r'),
           'x', ('<t', 'sub'), (', ', 'r'), 'y', ('<t', 'sub'), (')', 'r')), 'x'),
        (7.5, '2  condition the follower', ['qx', 'qm'],
         m('y', ('t', 'sub'), (' ∼ ', 'r'), 'p', ('(', 'r'), 'y', ('t', 'sub'), (' | ', 'r'),
           'x', ('t', 'sub'), (', ', 'r'), 'x', ('<t', 'sub'), (', ', 'r'), 'y', ('<t', 'sub'),
           (')', 'r')), 'y'),
        (10.9, '3  commit, swap roles', ['x', 'y'], None, None),
    ]
    for x0, title, qk, formula, out in steps:
        S.text(x0, 0.36, 3.2, 0.25, [(title, 'b')], size=10, align='l')
        y = 1.3
        labels = prefix_labels + ([m('q', ('x', 'sub')), m('q', ('y', 'sub'))] if qk[0] != 'x'
                                  else [m('x', ('t', 'sub')), m('y', ('t', 'sub'))])
        kinds = prefix_kinds + qk
        xs, xe = token_row(x0, y, labels, kinds)
        # committed prefix bracket
        S.line(x0, y + tw + 0.24, xs[3] + tw / 2, y + tw + 0.24, color=INK_2, lw=0.5)
        S.text(x0, y + tw + 0.26, xs[3] + tw / 2 - x0, 0.18, plain('committed prefix'), size=7,
               color=INK_2)
        if formula:
            src = xs[3] if out == 'x' else xs[5]
            fill, line = (BLUE_L, BLUE) if out == 'x' else (RED_L, RED)
            S.line(src, y - 0.04, src, y - 0.28, arrow=True, lw=0.75)
            S.rect(src - tw / 2, y - 0.6, tw, tw, fill=fill, line=line, lw=1.0)
            S.text(src + tw / 2 + 0.06, y - 0.62, 2.2, 0.34, formula, size=9, align='l')
            if out == 'y':                       # within-frame link q_x -> q_y
                S.line(xs[4] + tw / 2 - 0.05, y - 0.08, xs[5] - tw / 2 + 0.02, y - 0.08,
                       color=GOLD, lw=1.2, arrow=True)
        else:
            S.text(xe + 0.05, y - 0.02, 1.3, 0.36, [('→ at ', ''), ('t', 'i'), ('+1, ', ''),
                                                    ('y', 'i'), (' leads', '')], size=8,
                   align='l')

    # ------------------------------------------------ panel A: block (left)
    S.text(0.4, 1.4, 3.0, 0.3, [('A', 'b'), ('  Model architecture', '')], size=12, align='l')
    n = 3
    cols = []                                     # x centres of the sequence columns
    x = 0.85
    for i in range(n):
        cols += [x, x + 0.65]; x += 1.3
    x_dots = x + 0.05; x += 0.5
    cols += [x, x + 0.65]; x += 1.3              # x_T, y_T
    x += 0.35
    cols += [x, x + 0.65]                          # harmonizers
    tw2 = 0.42
    y_in, y_out = 7.72, 2.2
    lab_in = [m('x', ('0', 'sub')), m('y', ('0', 'sub')), m('x', ('1', 'sub')), m('y', ('1', 'sub')),
              m('x', ('2', 'sub')), m('y', ('2', 'sub')), m('x', ('T', 'sub')), m('y', ('T', 'sub')),
              m('q', ('x', 'sub')), m('q', ('y', 'sub'))]
    lab_out = [m('x', ('1', 'sub')), m('y', ('1', 'sub')), m('x', ('2', 'sub')), m('y', ('2', 'sub')),
               m('x', ('3', 'sub')), m('y', ('3', 'sub')), m('x', ('T+1', 'sub')),
               m('y', ('T+1', 'sub')), m('x', ('T+1', 'sub')), m('y', ('T+1', 'sub'))]
    for i, cx in enumerate(cols):
        kind = 'q' if i >= 8 else ('x' if i % 2 == 0 else 'y')
        fill, line = {'x': (BLUE_L, BLUE), 'y': (RED_L, RED), 'q': (GOLD_L, GOLD)}[kind]
        S.rect(cx - tw2 / 2, y_in, tw2, tw2, fill=fill, line=line, lw=1.0)
        S.text(cx - 0.3, y_in + tw2 + 0.02, 0.6, 0.22, lab_in[i], size=9)
        S.rect(cx - tw2 / 2, y_out, tw2, tw2, fill=fill, line=line, lw=1.0)
        S.text(cx - 0.3, y_out - 0.25, 0.6, 0.22, lab_out[i], size=9)
    S.text(x_dots - 0.15, y_in + 0.05, 0.4, 0.3, plain('…'), size=12)
    S.text(x_dots - 0.15, y_out + 0.05, 0.4, 0.3, plain('…'), size=12)
    S.text(cols[0] - 0.3, y_in + tw2 + 0.26, 6.0, 0.18, [('content tokens, shifted by one frame (', ''), ('x', 'i'), ('0', 'sub i'), (', ', ''), ('y', 'i'), ('0', 'sub i'), (': start)', '')],
           size=7.5, color=INK_2, align='l')
    S.text(cols[8] - 0.35, y_in + tw2 + 0.26, 1.4, 0.18, plain('harmonizers'), size=7.5, color=INK_2,
           align='l')
    S.text(cols[0] - 0.3, y_out - 0.47, 3.0, 0.2, plain('predicted tokens'), size=7.5, color=INK_2,
           align='l')

    # block frame
    bx, by, bw, bh = 0.55, 2.7, 8.65, 4.75
    S.rect(bx, by, bw, bh, fill='FFFFFF', line=INK, lw=1.0, z=0)
    S.text(bx + 0.15, by + 0.04, 3.0, 0.28, [('Transformer block ×', ''), ('L', 'i')], size=11,
           align='l')
    S.line(cols[4], y_in - 0.02, cols[4], by + bh + 0.02, arrow=True, lw=1.0)
    S.line(cols[4], by - 0.02, cols[4], y_out + tw2 + 0.02, arrow=True, lw=1.0)

    # spine (residual stream) on the right, post-LN: add, then LN
    xs_ = bx + bw - 0.45
    y_add_att, y_add_ffn = 5.45, 3.3
    S.line(xs_, by + bh - 0.35, xs_, by + 0.3, lw=0.75, arrow=True)
    for y_add in (y_add_att, y_add_ffn):
        S.node(xs_, y_add)
        S.rect(xs_ - 0.22, y_add - 0.55, 0.44, 0.26, fill='FFFFFF', line=INK, lw=0.75,
               runs=plain('LN'), size=8, z=4)
    S.text(xs_ - 0.35, by + bh - 0.3, 0.7, 0.2, plain('residual'), size=7, color=INK_2)

    # MoE FFN sublayer (top of the block)
    y_ffn = 3.8
    pi_x = (0.61, 0.08, 0.27, 0.04)
    pi_y = (0.04, 0.27, 0.61, 0.08)
    S.rect(bx + 0.2, 3.05, bw - 1.05, 1.95, fill=PANEL, line=None, z=0)
    S.text(bx + 0.3, 3.1, 1.2, 0.25, plain('MoE FFN'), size=9.5, align='l')
    ffn_x = [0.95, 2.85, 4.75, 6.65]
    ffn_w = 1.45
    for i, fx in enumerate(ffn_x):
        src = 'x' if i < 2 else 'y'
        fill = BLUE_L if src == 'x' else RED_L
        S.rect(fx, y_ffn, ffn_w, 0.62, fill=fill, line=INK, lw=0.75, dash=True,
               runs=[('FFN', ''), (str(i + 1), 'sub'), ('  (from LM', ''), (src, 'sub i'), (')', ''),
                     ('\n', ''), ('π', 'i'), ('x', 'sub i'), (f' = {pi_x[i]:.2f}   ', ''),
                     ('π', 'i'), ('y', 'sub i'), (f' = {pi_y[i]:.2f}', '')], size=8.5)
    # per-stream weighted sums above the experts (selected experts only)
    for (cx_sum, pis, col, rx) in ((2.6, pi_x, BLUE, 2.3), (6.3, pi_y, RED, 5.4)):
        S.node(cx_sum, 3.38)
        for fx, p in zip(ffn_x, pis):
            if p < 0.27:
                continue
            S.line(fx + ffn_w / 2 + (0.15 if col == RED else -0.15), y_ffn - 0.02,
                   cx_sum + (-0.08 if fx + ffn_w / 2 < cx_sum else 0.08), 3.5,
                   color=col, lw=1.0, arrow=True)
        S.line(cx_sum, 3.27, cx_sum, 3.12, lw=0.75, arrow=False)
    # routers under the experts, fan-out with pi labels
    y_r = 4.66
    for (rx, col, fill, nm, pis) in ((2.2, BLUE, BLUE_L, 'x', pi_x), (5.3, RED, RED_L, 'y', pi_y)):
        S.rect(rx, y_r, 1.3, 0.32, fill=fill, line=col, lw=0.75,
               runs=[('router ', ''), ('G', 'i'), ('(' + nm + ')', 'sup i')], size=8.5)
        for fx, p in zip(ffn_x, pis):
            on = p >= 0.27
            S.line(rx + 0.65, y_r - 0.02, fx + ffn_w / 2 + (0.15 if col == RED else -0.15),
                   y_ffn + 0.64, color=col if on else '999999', lw=1.0 if on else 0.5,
                   dash=not on, arrow=True)
    S.line(bx + bw - 1.05, 3.38, xs_ - 0.12, 3.38, lw=0.75)   # both sums -> residual add
    S.line(6.42, 3.38, bx + bw - 1.05, 3.38, lw=0.75)
    S.line(2.72, 3.38, 6.18, 3.38, lw=0.75)

    # attention sublayer
    y_wo = 5.3
    S.rect(bx + 0.2, 5.15, bw - 1.05, 2.15, fill=PANEL, line=None, z=0)
    S.text(bx + 0.3, 5.2, 2.6, 0.25, [('Duet attention  (masks: ', ''), ('B', 'b'), (')', '')], size=9.5,
           align='l')
    S.rect(3.3, y_wo, 3.0, 0.32, fill='FFFFFF', line=INK, lw=0.75, dash=True,
           runs=[('output projection  ', ''), ('W', 'i'), ('O', 'sup i'), ('x', 'sub i'), (' / ', ''),
                 ('W', 'i'), ('O', 'sup i'), ('y', 'sub i')], size=8.5)
    S.line(6.32, y_wo + 0.16, xs_ - 0.12, y_wo + 0.16, lw=0.75)   # -> residual add
    cx_a = 4.8
    S.node(cx_a, 5.85)
    S.line(cx_a, 5.73, cx_a, y_wo + 0.34, arrow=True, lw=0.75)
    # three passes
    passes = [(0.85, 'intra-stream', 'K', 'intra', BLUE_L, BLUE, None),
              (3.5, 'cross-stream', 'K', 'cross', RED_L, RED, ('g', 'c')),
              (6.15, 'within-frame', 'K', 'frame', GOLD_L, GOLD, ('g', 'f'))]
    y_p = 6.3
    for px_, name, K, sub, fill, line, gate in passes:
        S.rect(px_, y_p, 2.35, 0.5, fill=fill, line=line, lw=0.75,
               runs=[(name + '  ', ''), ('K', 'i'), (sub, 'sub')], size=8.5)
        xm = px_ + 2.35 / 2
        if gate:
            S.rect(xm - 0.15, 5.95, 0.3, 0.3, fill=GOLD_L, line=GOLD, lw=0.75, shape='ellipse',
                   runs=[('g', 'i'), (gate[1], 'sup i')], size=8, z=4)
            S.line(xm, y_p - 0.02, xm, 6.27, arrow=True, lw=0.75)
            S.line(xm, 5.93, cx_a + (0.1 if xm > cx_a else -0.1), 5.93, arrow=True, lw=0.75)
        else:
            S.line(xm, y_p - 0.02, xm, 5.85, lw=0.75)
            S.line(xm, 5.85, cx_a - 0.12, 5.85, arrow=True, lw=0.75)
    # per-stream projections
    y_w = 6.95
    for (wx, s, fill) in ((1.05, 'x', BLUE_L), (4.95, 'y', RED_L)):
        S.rect(wx, y_w, 2.9, 0.32, fill=fill, line=INK, lw=0.75, dash=True,
               runs=[('W', 'i'), ('Q', 'sup i'), (s, 'sub i'), (',  ', ''),
                     ('W', 'i'), ('K', 'sup i'), (s, 'sub i'), (',  ', ''),
                     ('W', 'i'), ('V', 'sup i'), (s, 'sub i')], size=9)

    # legend (bottom-left)
    lx, ly = 0.65, 8.62
    for k, (fill, line, dash, lab) in enumerate((
            (BLUE_L, BLUE, False, [('stream ', ''), ('x', 'i'), (' (melody)', '')]),
            (BLUE_L, INK, True, [('weights from LM', ''), ('x', 'sub i')]),
            (RED_L, RED, False, [('stream ', ''), ('y', 'i'), (' (chord)', '')]),
            (RED_L, INK, True, [('weights from LM', ''), ('y', 'sub i')]),
            (GOLD_L, GOLD, False, plain('harmonizer')),
            ('FFFFFF', GOLD, True, plain('masked harmonizer')))):
        xx = lx + k * 1.6
        S.rect(xx, ly, 0.28, 0.2, fill=fill, line=line, lw=0.75, dash=dash)
        S.text(xx + 0.34, ly - 0.03, 1.3, 0.26, lab, size=8, align='l')

    # ------------------------------------------------ panel B: admission map
    S.text(9.95, 2.85, 4.0, 0.3, [('B', 'b'), ('  Attention design', '')], size=12, align='l')
    from cp_transformer_m2c_duet_block import M2CDuetBlockLayer
    layer = M2CDuetBlockLayer(hidden_size=16, num_heads=2, intermediate_size=32,
                              moe_num_experts=1, moe_topk=1, moe_intermediate_size=32)
    clean_len = 2 * T_show
    intra, cross, frame = layer._build_masks(clean_len, T_show, torch.device('cpu'))
    diag = np.eye(clean_len + 2, dtype=bool)
    cross = cross.numpy() & ~diag
    intra, frame = intra.numpy(), frame.numpy()
    L = clean_len + 2
    labels = [m('x', (str(i // 2 + 1), 'sub')) if i % 2 == 0 else m('y', (str(i // 2 + 1), 'sub'))
              for i in range(clean_len)] + [m('q', ('x', 'sub')), m('q', ('y', 'sub'))]
    cs = 0.31
    ox, oy = 10.5, 3.72
    S.text(ox, oy - 0.5, L * cs, 0.2, plain('keys'), size=8, color=INK_2)
    S.text(ox - 0.85, oy + L * cs / 2 - 0.1, 0.5, 0.2, plain('queries'), size=8, color=INK_2)
    for i in range(L):
        S.text(ox + i * cs - 0.1, oy - 0.28, cs + 0.2, 0.22, labels[i], size=8,
               color=GOLD if i >= clean_len else (BLUE if i % 2 == 0 else RED))
        S.text(ox - 0.42, oy + i * cs + 0.04, 0.35, 0.22, labels[i], size=8, align='r',
               color=GOLD if i >= clean_len else (BLUE if i % 2 == 0 else RED))
        for j in range(L):
            key = 'intra' if intra[i, j] else 'cross' if cross[i, j] else 'frame' if frame[i, j] else None
            if key:
                fill, line = {'intra': (BLUE_L, BLUE), 'cross': (RED_L, RED), 'frame': (GOLD_L, GOLD)}[key]
                S.rect(ox + j * cs, oy + i * cs, cs, cs, fill=fill, line=line, lw=0.5)
            else:
                blocked = j > i and i < clean_len
                S.rect(ox + j * cs, oy + i * cs, cs, cs, fill='F4F4F4' if blocked else 'FFFFFF',
                       line='DBD9D4', lw=0.4)
    S.rect(ox, oy, L * cs, L * cs, fill=None, line=INK, lw=1.0, z=3)
    S.line(ox, oy + clean_len * cs, ox + L * cs, oy + clean_len * cs, lw=1.0, z=3)
    S.line(ox + clean_len * cs, oy, ox + clean_len * cs, oy + L * cs, lw=1.0, z=3)
    for k, (fill, line, lab) in enumerate(((BLUE_L, BLUE, [('intra-stream  ', ''), ('K', 'i'), ('intra', 'sub')]),
                                          (RED_L, RED, [('cross-stream  ', ''), ('K', 'i'), ('cross', 'sub'), ('  × ', ''), ('g', 'i'), ('c', 'sup i')]),
                                          (GOLD_L, GOLD, [('within-frame  ', ''), ('K', 'i'), ('frame', 'sub'), ('  × ', ''), ('g', 'i'), ('f', 'sup i')]),
                                          ('F4F4F4', 'DBD9D4', plain('blocked by causality')))):
        xx, yy = ox + (k % 2) * 2.0, oy + L * cs + 0.35 + (k // 2) * 0.32
        S.rect(xx, yy, 0.28, 0.2, fill=fill, line=line, lw=0.6)
        S.text(xx + 0.34, yy - 0.03, 1.7, 0.26, lab, size=8, align='l')
    return S


# ---------------------------------------------------------------- pptx backend
def write_pptx(S, path):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.enum.dml import MSO_LINE
    from pptx.oxml.ns import qn
    from lxml import etree

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SLIDE_W), Inches(SLIDE_H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    def rgb(h):
        return RGBColor.from_string(h)

    def set_runs(tf, runs, size, color, align, bold=False, valign='m'):
        tf.word_wrap = False
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = {'m': MSO_ANCHOR.MIDDLE, 't': MSO_ANCHOR.TOP}[valign]
        p = tf.paragraphs[0]
        p.alignment = {'c': PP_ALIGN.CENTER, 'l': PP_ALIGN.LEFT, 'r': PP_ALIGN.RIGHT}[align]
        for text, flags in runs:
            if text == '\n':
                p = tf.add_paragraph()
                p.alignment = {'c': PP_ALIGN.CENTER, 'l': PP_ALIGN.LEFT, 'r': PP_ALIGN.RIGHT}[align]
                continue
            r = p.add_run(); r.text = text
            f = r.font; f.name = FONT; f.size = Pt(size); f.color.rgb = rgb(color)
            f.italic = 'i' in flags.split()
            f.bold = bold or 'b' in flags.split()
            if 'sup' in flags.split():
                r.font._element.set('baseline', '30000')
            if 'sub' in flags.split():
                r.font._element.set('baseline', '-25000')

    for it in sorted(S.items, key=lambda d: d['z']):
        if it['k'] in ('rect',):
            shp = {'rect': MSO_SHAPE.RECTANGLE, 'ellipse': MSO_SHAPE.OVAL}[it['shape']]
            s = slide.shapes.add_shape(shp, Inches(it['x']), Inches(it['y']), Inches(it['w']),
                                       Inches(it['h']))
            s.shadow.inherit = False
            if it['fill'] is None:
                s.fill.background()
            else:
                s.fill.solid(); s.fill.fore_color.rgb = rgb(it['fill'])
            if it['line'] is None:
                s.line.fill.background()
            else:
                s.line.color.rgb = rgb(it['line']); s.line.width = Pt(it['lw'])
                if it['dash']:
                    s.line.dash_style = MSO_LINE.SQUARE_DOT
            if it['runs']:
                set_runs(s.text_frame, it['runs'], it['size'], it['color'], it['align'])
            else:
                s.text_frame.text = ''
        elif it['k'] == 'text':
            tb = slide.shapes.add_textbox(Inches(it['x']), Inches(it['y']), Inches(it['w']),
                                          Inches(it['h']))
            set_runs(tb.text_frame, it['runs'], it['size'], it['color'], it['align'],
                     bold=it['bold'], valign=it['valign'])
        elif it['k'] == 'image':
            slide.shapes.add_picture(it['path'], Inches(it['x']), Inches(it['y']),
                                     Inches(it['w']), Inches(it['h']))
        elif it['k'] == 'line':
            c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(it['x1']), Inches(it['y1']),
                                           Inches(it['x2']), Inches(it['y2']))
            c.line.color.rgb = rgb(it['color']); c.line.width = Pt(it['lw'])
            if it['dash']:
                c.line.dash_style = MSO_LINE.DASH
            if it['arrow']:
                ln = c.line._get_or_add_ln()
                tail = etree.SubElement(ln, qn('a:tailEnd'))
                tail.set('type', 'triangle'); tail.set('w', 'med'); tail.set('len', 'med')
    prs.save(path)


def write_pptx_js(S, path):
    """Build the .pptx with pptxgenjs (spec_to_pptx.js) instead of
    python-pptx; PowerPoint opens pptxgenjs output reliably."""
    import json
    import subprocess
    here = os.path.dirname(os.path.abspath(__file__))
    spec = {'width': SLIDE_W, 'height': SLIDE_H, 'font': FONT, 'items': S.items}
    tmp = path + '.spec.json'
    with open(tmp, 'w') as f:
        json.dump(spec, f)
    env = dict(os.environ)
    # pptxgenjs must be resolvable: `npm install pptxgenjs` next to this
    # file, or point NODE_PATH at a node_modules that has it
    local = os.path.join(here, 'node_modules')
    if os.path.isdir(os.path.join(local, 'pptxgenjs')):
        env['NODE_PATH'] = local + os.pathsep + env.get('NODE_PATH', '')
    subprocess.run(['node', os.path.join(here, 'spec_to_pptx.js'), tmp, path], check=True, env=env)
    os.remove(tmp)


# ---------------------------------------------------------------- preview backend
def runs_to_mathtext(runs):
    out = ''
    for text, flags in runs:
        fl = flags.split()
        t = text.replace('−', '-').replace('∼', r'\sim').replace('π', r'\pi').replace('𝒦', r'\mathcal{K}')
        t = t.replace('→', r'\rightarrow').replace('…', r'\ldots')
        if 'sup' in fl:
            out += '$^{' + (t if 'i' in fl else r'\mathrm{' + t + '}') + '}$'
        elif 'sub' in fl:
            out += '$_{' + (t if 'i' in fl else r'\mathrm{' + t + '}') + '}$'
        elif 'i' in fl:
            out += '$' + t.replace(' ', r'\ ') + '$'
        elif text == '\n':
            out += '\n'
        else:
            out += text
    return out


def write_preview(S, path):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib.patches import Rectangle, Ellipse, FancyArrowPatch
    fig = plt.figure(figsize=(SLIDE_W, SLIDE_H))
    ax = fig.add_axes([0, 0, 1, 1]); ax.set_xlim(0, SLIDE_W); ax.set_ylim(SLIDE_H, 0); ax.axis('off')
    hexc = lambda h: '#' + h
    for it in sorted(S.items, key=lambda d: d['z']):
        if it['k'] == 'rect':
            kw = dict(fc=hexc(it['fill']) if it['fill'] else 'none',
                      ec=hexc(it['line']) if it['line'] else 'none', lw=it['lw'],
                      ls=(0, (1.5, 1.5)) if it['dash'] else '-', zorder=it['z'])
            if it['shape'] == 'ellipse':
                ax.add_patch(Ellipse((it['x'] + it['w'] / 2, it['y'] + it['h'] / 2), it['w'], it['h'], **kw))
            else:
                ax.add_patch(Rectangle((it['x'], it['y']), it['w'], it['h'], **kw))
            if it['runs']:
                ax.text(it['x'] + it['w'] / 2, it['y'] + it['h'] / 2, runs_to_mathtext(it['runs']),
                        ha='center', va='center', fontsize=it['size'], color=hexc(it['color']),
                        zorder=it['z'] + 0.5, family='serif')
        elif it['k'] == 'text':
            ha = {'c': 'center', 'l': 'left', 'r': 'right'}[it['align']]
            xx = {'c': it['x'] + it['w'] / 2, 'l': it['x'], 'r': it['x'] + it['w']}[it['align']]
            ax.text(xx, it['y'] + it['h'] / 2, runs_to_mathtext(it['runs']), ha=ha, va='center',
                    fontsize=it['size'], color=hexc(it['color']), zorder=it['z'], family='serif',
                    weight='bold' if it['bold'] else 'normal')
        elif it['k'] == 'image':
            img = plt.imread(it['path'])
            ax.imshow(img, extent=(it['x'], it['x'] + it['w'], it['y'] + it['h'], it['y']),
                      zorder=it['z'], aspect='auto', interpolation='bilinear')
        elif it['k'] == 'poly':
            from matplotlib.patches import Polygon
            ax.add_patch(Polygon(it['points'], closed=True, fc=hexc(it['fill']), ec='none',
                                 alpha=it['alpha'], zorder=it['z']))
        elif it['k'] == 'curve':
            from matplotlib.path import Path
            dy = (it['y2'] - it['y1']) * it['bend']
            verts = [(it['x1'], it['y1']), (it['x1'], it['y1'] + dy), (it['x2'], it['y2'] - dy),
                     (it['x2'], it['y2'])]
            a = FancyArrowPatch(path=Path(verts, [Path.MOVETO, Path.CURVE4, Path.CURVE4,
                                                  Path.CURVE4]),
                                arrowstyle='-|>' if it['arrow'] else '-', mutation_scale=7,
                                color=hexc(it['color']), lw=it['lw'], zorder=it['z'], fill=False)
            ax.add_patch(a)
        elif it['k'] == 'line':
            a = FancyArrowPatch((it['x1'], it['y1']), (it['x2'], it['y2']),
                                arrowstyle='-|>' if it['arrow'] else '-', mutation_scale=7,
                                color=hexc(it['color']), lw=it['lw'], zorder=it['z'],
                                shrinkA=0, shrinkB=0)
            if it['dash']:
                a.set_linestyle((0, (3, 2)))
            ax.add_patch(a)
    fig.savefig(path, dpi=150)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--out', required=True, help='path without extension')
    args = ap.parse_args()
    S = build()
    write_pptx(S, args.out + '.pptx'); print('wrote', args.out + '.pptx')
    write_preview(S, args.out + '.png'); print('wrote', args.out + '.png')


if __name__ == '__main__':
    main()
